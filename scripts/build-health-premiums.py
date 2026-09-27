#!/usr/bin/env python3
"""
Build lib/data/health-premiums.json — mandatory basic health insurance (OKP)
premiums per premium region, plus the BFS commune → premium region mapping.

Run:  python3 scripts/build-health-premiums.py
Needs openpyxl (pip install openpyxl) for FOPH's region workbook.

Sources (FOPH / BAG, published once a year, late September, for the next year):
  - Premiums: dataset "Krankenversicherungsprämien" on opendata.swiss
    (https://opendata.swiss/de/dataset/health-insurance-premiums), file
    Prämien_CH.csv. Every approved premium: insurer x canton x region x age
    class x accident cover x model x deductible.
  - Regions: FOPH's "Prämienregionen" workbook on priminfo
    (https://www.priminfo.admin.ch/downloads/praemienregionen.xlsx), sheet
    A_COM: BFS commune number → canton → premium region. The BFS number, not
    the postcode, decides the region.

Which year: the newest year with published data. The script prefers the
current Prämien_CH.csv and falls back to the newest yearly archive when the
current file is an empty placeholder (FOPH uploads header-only files ahead of
each release). NOTE: FOPH announced a new coding for the 2027 files
(PR_REG_x regions, AKA_* age classes, BASE/PRAXIS/... models, new deductible
codes); the code constants below are the 2026 scheme and the script exits on
unknown codes rather than guessing, so update them when 2027 is published.

Averaging method (MVP), per canton + premium region + age class:
  - model: standard (Tariftyp TAR-BASE — free choice of doctor)
  - deductible: the ordinary one (isBaseF = 1): CHF 300 for adults and young
    adults, CHF 0 for children
  - accident cover included (MIT-UNF). Employees working 8+ hours a week are
    covered through their employer and can drop it for a lower premium; the
    with-accident figure is the conservative, FOPH-reference variant.
  - children: the regular child premium (sub-group K1). Some insurers file a
    cheaper tier for the 3rd and later child (K3/K4/K5); not applied here.
  - unweighted arithmetic mean across all insurers offering that premium.
    FOPH's own "mittlere Prämie" is weighted by insured population across all
    models and deductibles, so it is lower and not directly comparable.
"""

import csv
import io
import json
import re
import statistics
import sys
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from base64 import b64encode
from collections import defaultdict
from datetime import date
from pathlib import Path

import openpyxl

UA = 'Mozilla/5.0 (compatible; liveability-data-import)'
BAG = 'https://opendata.bagnet.ch/?r=/download&path='
REGIONS_URL = 'https://www.priminfo.admin.ch/downloads/praemienregionen.xlsx'
AGV = 'https://www.agvchapp.bfs.admin.ch/api/communes/snapshot?date=01-01-{year}'
OUT = Path(__file__).resolve().parent.parent / 'lib' / 'data' / 'health-premiums.json'

CANTONS = {'AG', 'AI', 'AR', 'BE', 'BL', 'BS', 'FR', 'GE', 'GL', 'GR', 'JU', 'LU', 'NE',
           'NW', 'OW', 'SG', 'SH', 'SO', 'SZ', 'TG', 'TI', 'UR', 'VD', 'VS', 'ZG', 'ZH'}

# 2026 coding scheme (see "Erläuterungen zu den Prämiendaten")
AGE_CLASSES = {'AKL-KIN': 'child', 'AKL-JUG': 'youngAdult', 'AKL-ERW': 'adult'}
STANDARD_MODEL = 'TAR-BASE'
WITH_ACCIDENT = 'MIT-UNF'
REGULAR_CHILD_TIER = 'K1'
REGION_CODE = re.compile(r'^PR-REG CH(\d)$')

# Communes formed after FOPH's region workbook was cut: new BFS → predecessors.
# The predecessors must all share one region, which the new commune inherits.
MERGERS = {
    2056: (2016, 2027),  # Fétigny-Ménières (FR), 1.1.2026 ← Fétigny + Ménières
}


def download(url: str) -> bytes:
    req = urllib.request.Request(url, headers={'User-Agent': UA})
    with urllib.request.urlopen(req, timeout=300) as res:
        return res.read()


def bag_url(path: str) -> str:
    return BAG + urllib.parse.quote(b64encode(path.encode()).decode())


def load_premium_rows() -> list[dict]:
    current = download(bag_url('/Praemien/Prämien_CH.csv')).decode('utf-8-sig')
    rows = list(csv.DictReader(io.StringIO(current)))
    if rows:
        print(f'Using current Prämien_CH.csv ({len(rows)} rows)')
        return rows
    # Placeholder: take the newest archive that exists
    for year in range(date.today().year + 1, date.today().year - 2, -1):
        try:
            archive = download(bag_url(f'/Praemien/Archiv_Praemien_{year}.zip'))
            z = zipfile.ZipFile(io.BytesIO(archive))
        except (urllib.error.HTTPError, zipfile.BadZipFile):
            continue
        member = next(i for i in z.infolist() if i.filename.endswith('_CH.csv') and 'mien' in i.filename)
        rows = list(csv.DictReader(io.StringIO(z.read(member).decode('utf-8-sig'))))
        print(f'Current Prämien_CH.csv is an empty placeholder; using Archiv_Praemien_{year}.zip ({len(rows)} rows)')
        return rows
    sys.exit('No premium data found')


def regional_averages(rows: list[dict]) -> tuple[int, dict]:
    years = {r['Geschäftsjahr'] for r in rows}
    if len(years) != 1:
        sys.exit(f'Expected one premium year, got {years}')
    premium_year = int(years.pop())

    unknown_age = {r['Altersklasse'] for r in rows} - set(AGE_CLASSES)
    unknown_region = {r['Region'] for r in rows if not REGION_CODE.match(r['Region'])}
    if unknown_age or unknown_region or STANDARD_MODEL not in {r['Tariftyp'] for r in rows}:
        sys.exit(f'Unknown codes (new FOPH scheme?) age={unknown_age} region={unknown_region}; update constants')

    samples: dict[tuple[str, int, str], dict[str, float]] = defaultdict(dict)
    for r in rows:
        if (r['Kanton'] in CANTONS and r['Tariftyp'] == STANDARD_MODEL and r['isBaseF'] == '1'
                and r['Unfalleinschluss'] == WITH_ACCIDENT and r['Altersuntergruppe'] in ('', REGULAR_CHILD_TIER)):
            key = (r['Kanton'], int(REGION_CODE.match(r['Region']).group(1)), AGE_CLASSES[r['Altersklasse']])
            if r['Versicherer'] in samples[key]:
                sys.exit(f'Insurer {r["Versicherer"]} has two standard premiums for {key}')
            samples[key][r['Versicherer']] = float(r['Prämie'])

    regions: dict[str, dict] = defaultdict(dict)
    for (canton, region, age), by_insurer in sorted(samples.items()):
        regions[f'{canton}-{region}'][age] = round(statistics.fmean(by_insurer.values()), 2)
        regions[f'{canton}-{region}'].setdefault('insurers', {})[age] = len(by_insurer)
    for key, value in regions.items():
        missing = set(AGE_CLASSES.values()) - set(value)
        if missing:
            sys.exit(f'Region {key} lacks {missing}')
    return premium_year, dict(regions)


def commune_regions(premium_year: int) -> dict[int, str]:
    wb = openpyxl.load_workbook(io.BytesIO(download(REGIONS_URL)), read_only=True)
    title = str(next(wb['Informationen'].iter_rows(values_only=True))[0])
    if f'31.12.{premium_year}' not in title:
        sys.exit(f'Region workbook is for another year: {title!r}')

    communes: dict[int, str] = {}
    for row in wb['A_COM'].iter_rows(min_row=6, values_only=True):
        if isinstance(row[0], int):
            key = f'{row[1]}-{int(row[3])}'
            if communes.setdefault(row[0], key) != key:
                sys.exit(f'BFS {row[0]} maps to two regions')

    for new_bfs, predecessors in MERGERS.items():
        if new_bfs in communes:
            continue
        keys = {communes[p] for p in predecessors}
        if len(keys) != 1:
            sys.exit(f'Merged commune {new_bfs} predecessors disagree: {keys}')
        communes[new_bfs] = keys.pop()

    # Keep exactly the communes that exist on 1 January of the premium year
    # (the workbook still lists some retired BFS numbers), and fail on gaps.
    agv = download(AGV.format(year=premium_year)).decode('utf-8-sig')
    fso = {int(r['BfsCode']) for r in csv.DictReader(io.StringIO(agv)) if r['Level'] == '3'}
    missing = sorted(fso - set(communes))
    if missing:
        sys.exit(f'No premium region for BFS {missing}; add them to MERGERS')
    return {b: communes[b] for b in fso}


def main() -> None:
    premium_year, regions = regional_averages(load_premium_rows())
    communes = commune_regions(premium_year)

    # Every commune must land in a region that has averages, and vice versa
    unused = set(regions) - set(communes.values())
    orphans = {b: k for b, k in communes.items() if k not in regions}
    if unused or orphans:
        sys.exit(f'Region mismatch — no communes: {unused}; communes without premiums: {orphans}')

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({
        'premiumYear': premium_year,
        'generated': date.today().isoformat(),
        'source': 'FOPH (BAG): Prämien_CH (opendata.swiss health-insurance-premiums) and '
                  'Prämienregionen (priminfo.admin.ch)',
        'method': 'Unweighted mean across insurers of the monthly premium for the standard model '
                  'with the ordinary deductible (CHF 300 adults and young adults, CHF 0 children), '
                  'accident cover included, regular child premium (no 3rd-child discount).',
        'regions': regions,
        'communes': {str(b): k for b, k in sorted(communes.items())},
    }, ensure_ascii=False, separators=(',', ':')) + '\n')
    print(f'Premium year {premium_year}: {len(regions)} regions, {len(communes)} communes → '
          f'{OUT.relative_to(OUT.parent.parent.parent)}')


if __name__ == '__main__':
    main()
